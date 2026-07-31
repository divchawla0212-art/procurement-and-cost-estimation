"""Create a tiny in-memory-like BOQ workbook for deterministic tests."""
import openpyxl


def make(path: str) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sect. 3 D - Electrical"
    ws["B7"] = "Item"; ws["C7"] = "Description"; ws["E7"] = "UoM"
    ws["F7"] = "Total Quantity (Q)"; ws["K7"] = "Unit Price"; ws["L7"] = "TOTAL"
    # priced row
    ws["B8"] = "E.05.01"; ws["C8"] = "LV cable"; ws["E8"] = "m"
    ws["F8"] = 10; ws["K8"] = 5.0; ws["L8"] = 50.0
    # blank-priced row (qty only)
    ws["B9"] = "E.05.02"; ws["C9"] = "HV cable"; ws["E9"] = "m"
    ws["F9"] = 20; ws["K9"] = 0; ws["L9"] = 0
    # non-item row (header/section, no code)
    ws["C10"] = "SUBTOTAL"
    wb.save(path)
    return path
