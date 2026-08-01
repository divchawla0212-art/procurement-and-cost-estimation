"""Regenerate `rfq_clauses.docx`, the committed Word fixture.

The fixture is committed as a binary rather than built during the test run on
purpose. Building it in-process would read back only what python-docx had just
written, so a reader that understood nothing but its own output would still pass.
The committed file is a real OOXML package, which is the thing `read_docx_text`
actually has to parse.

Run from the repo root after editing:

    python tests/fixtures/make_rfq_docx.py

Clause text lives in *both* a paragraph and a table on purpose: RFQ clauses
frequently sit in tables, and a reader that walked paragraphs alone would drop
them silently. `Rated output` appears only inside the table for exactly that
reason - it is the token that fails when tables are skipped.
"""
import docx


def make(path: str) -> str:
    doc = docx.Document()
    doc.add_paragraph("MATERIAL REQUISITION MR-4471")
    doc.add_paragraph("4.2.7 H2S tolerance shall be at least 50 ppm.")

    table = doc.add_table(rows=3, cols=3)
    rows = [
        ("Clause", "Requirement", "Unit"),
        ("5.1", "Rated output shall be at least 1500", "kVA"),
        ("5.2", "Ambient design temperature shall not exceed 50", "degC"),
    ]
    for row, values in zip(table.rows, rows):
        for cell, value in zip(row.cells, values):
            cell.text = value

    doc.add_paragraph("End of requisition.")
    doc.save(path)
    return path


if __name__ == "__main__":
    import os
    make(os.path.join(os.path.dirname(__file__), "rfq_clauses.docx"))
