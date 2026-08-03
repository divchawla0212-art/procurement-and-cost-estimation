import io
import openpyxl
from procurement.models import ComparisonRow, ComparisonTable
from procurement.matrix import ComplianceMatrix, MatrixCell, MatrixRow, bound
from procurement.export import (comparison_to_rows, comparison_to_xlsx_bytes,
                                comparison_to_csv_str, matrix_to_csv_str,
                                matrix_to_detail_rows, matrix_to_grid_rows,
                                matrix_to_xlsx_bytes)


def _table():
    return ComparisonTable(target_currency="USD", rows=[
        ComparisonRow(vendor="B", currency="USD", raw_base_price=1000.0, normalized_total=1000.0,
                      delivery_terms="CIF", delivery_time="24w", payment_terms="LC",
                      engine_make="MAN", extraction_status="ok"),
    ])


def test_rows_and_csv():
    rows = comparison_to_rows(_table())
    assert rows[0]["vendor"] == "B" and rows[0]["normalized_total"] == 1000.0
    csv = comparison_to_csv_str(_table())
    assert "vendor" in csv and "MAN" in csv


def test_xlsx_bytes_roundtrip():
    data = comparison_to_xlsx_bytes(_table())
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active
    assert ws["A1"].value == "vendor"
    assert any(cell.value == "B" for row in ws.iter_rows() for cell in row)


# --------------------------------------------------------- compliance matrix

def _matrix() -> ComplianceMatrix:
    """Two vendors, and deliberately only one cell on the auto row: the
    exporter's handling of a vendor the pipeline computed nothing for is the
    store invariant worth pinning, and `build_matrix` never produces it."""
    return ComplianceMatrix(
        vendors=["KERUI", "SUZLON"],
        rows=[
            MatrixRow(req_id="r-1", clause_ref="1.1", text="H2S at least 50 ppm",
                      checkability="auto", parameter="h2s", operator=">=",
                      value=50, unit="ppm",
                      cells={"KERUI": MatrixCell(
                          vendor="KERUI", verdict="fail", group="not_matched",
                          rationale="40 ppm >= 50 ppm is false",
                          fact_id="f-1", doc_id="d-9",
                          candidate_fact_ids=["f-1", "f-2"])}),
            MatrixRow(req_id="r-2", clause_ref="9.1",
                      text="Submit an O&M manual", checkability="judgement",
                      cells={"KERUI": MatrixCell(vendor="KERUI", verdict="review",
                                                 group="needs_human",
                                                 rationale="human judgement required"),
                             "SUZLON": MatrixCell(vendor="SUZLON", verdict="pass",
                                                  group="matched")}),
        ])


def test_grid_row_holds_the_bound_and_one_verdict_per_vendor():
    grid = matrix_to_grid_rows(_matrix())
    assert grid[0]["Clause"] == "1.1"
    assert grid[0]["Requirement"] == "h2s >= 50.0 ppm"
    assert grid[0]["KERUI"] == "fail"


def test_a_vendor_with_no_computed_cell_exports_blank_never_a_verdict():
    """INV: missing data is never coerced to a passing or zero value."""
    grid = matrix_to_grid_rows(_matrix())
    assert grid[0]["SUZLON"] == ""


def test_grid_carries_a_column_pair_for_every_vendor_in_matrix_order():
    grid = matrix_to_grid_rows(_matrix())
    assert list(grid[0]) == ["Clause", "Requirement", "Checkability",
                             "KERUI", "KERUI documents",
                             "SUZLON", "SUZLON documents"]


def test_detail_holds_one_row_per_requirement_and_vendor():
    detail = matrix_to_detail_rows(_matrix())
    assert len(detail) == 4
    assert [(r["Clause"], r["Vendor"]) for r in detail] == [
        ("1.1", "KERUI"), ("1.1", "SUZLON"),
        ("9.1", "KERUI"), ("9.1", "SUZLON")]


def test_detail_carries_the_rationale_verbatim_and_its_evidence():
    detail = matrix_to_detail_rows(_matrix())
    row = detail[0]
    assert row["Rationale"] == "40 ppm >= 50 ppm is false"
    assert row["Fact id"] == "f-1" and row["Doc id"] == "d-9"
    assert row["Candidate fact ids"] == "f-1, f-2"


def test_detail_keeps_the_clause_text_the_bound_replaces():
    """`Requirement` is the checkable bound for an auto row, so the prose the
    reviewer reads would otherwise be lost from the export entirely."""
    detail = matrix_to_detail_rows(_matrix())
    assert detail[0]["Clause text"] == "H2S at least 50 ppm"


def test_detail_leaves_an_absent_cell_blank_in_every_column():
    detail = matrix_to_detail_rows(_matrix())
    absent = detail[1]
    assert absent["Vendor"] == "SUZLON"
    assert absent["Verdict"] == "" and absent["Group"] == ""
    assert absent["Rationale"] == "" and absent["Fact id"] == ""


def test_matrix_csv_is_the_grid_sheet():
    csv = matrix_to_csv_str(_matrix())
    assert csv.splitlines()[0] == (
        "Clause,Requirement,Checkability,KERUI,KERUI documents,"
        "SUZLON,SUZLON documents")
    assert "human judgement required" not in csv


def test_matrix_xlsx_carries_both_the_grid_and_the_detail():
    wb = openpyxl.load_workbook(io.BytesIO(matrix_to_xlsx_bytes(_matrix())))
    assert wb.sheetnames == ["Matrix", "Detail"]
    assert wb["Matrix"]["A1"].value == "Clause"
    assert wb["Detail"]["A1"].value == "Clause"
    assert any(cell.value == "40 ppm >= 50 ppm is false"
               for row in wb["Detail"].iter_rows() for cell in row)


def test_an_empty_matrix_still_produces_a_readable_workbook():
    wb = openpyxl.load_workbook(io.BytesIO(matrix_to_xlsx_bytes(ComplianceMatrix())))
    assert wb.sheetnames == ["Matrix", "Detail"]
    assert matrix_to_csv_str(ComplianceMatrix()) == ""


def test_bound_renders_each_checkability_tier():
    rows = _matrix().rows
    assert bound(rows[0]) == "h2s >= 50.0 ppm"
    assert bound(rows[1]) == "Submit an O&M manual"
    stated = MatrixRow(req_id="r-3", clause_ref="2.1", text="State the make",
                       checkability="stated", parameter="engine_make")
    assert bound(stated) == "engine_make (must be stated)"


# ------------------------------------------- the document behind each verdict

def _named_matrix() -> ComplianceMatrix:
    return ComplianceMatrix(
        vendors=["KERUI", "SUZLON"],
        rows=[MatrixRow(
            req_id="r-1", clause_ref="1.1", text="H2S at least 50 ppm",
            checkability="auto", parameter="h2s", operator=">=", value=50,
            unit="ppm",
            cells={"KERUI": MatrixCell(
                       vendor="KERUI", verdict="fail", group="not_matched",
                       rationale="40 ppm >= 50 ppm is false",
                       fact_id="f-1", doc_id="d-9",
                       doc_name="datasheet.pdf",
                       candidate_doc_names=["datasheet.pdf"]),
                   "SUZLON": MatrixCell(
                       vendor="SUZLON", verdict="review", group="needs_human",
                       rationale="stated more than once",
                       fact_id="f-2", doc_id="d-7",
                       candidate_fact_ids=["f-2", "f-3"],
                       doc_name="offer.pdf",
                       candidate_doc_names=["offer.pdf", "annexure.pdf"])})])


def test_each_vendor_gets_a_documents_column_beside_its_verdict():
    grid = matrix_to_grid_rows(_named_matrix())
    assert list(grid[0]) == ["Clause", "Requirement", "Checkability",
                             "KERUI", "KERUI documents",
                             "SUZLON", "SUZLON documents"]


def test_the_documents_column_names_the_file_the_verdict_was_read_from():
    grid = matrix_to_grid_rows(_named_matrix())
    assert grid[0]["KERUI documents"] == "datasheet.pdf"


def test_a_verdict_read_from_two_documents_names_both():
    grid = matrix_to_grid_rows(_named_matrix())
    assert grid[0]["SUZLON documents"] == "offer.pdf; annexure.pdf"


def test_a_verdict_citing_no_document_leaves_its_documents_cell_blank():
    grid = matrix_to_grid_rows(_matrix())
    assert grid[1]["KERUI documents"] == ""
    assert grid[0]["SUZLON documents"] == ""


def test_the_csv_grid_carries_the_documents_columns_too():
    csv = matrix_to_csv_str(_named_matrix())
    assert csv.splitlines()[0] == (
        "Clause,Requirement,Checkability,KERUI,KERUI documents,"
        "SUZLON,SUZLON documents")


def test_detail_names_the_document_beside_its_opaque_id():
    detail = matrix_to_detail_rows(_named_matrix())
    assert detail[0]["Doc id"] == "d-9"
    assert detail[0]["Doc name"] == "datasheet.pdf"
    assert detail[1]["Candidate documents"] == "offer.pdf, annexure.pdf"


def test_the_named_grid_survives_a_workbook_roundtrip():
    wb = openpyxl.load_workbook(io.BytesIO(matrix_to_xlsx_bytes(_named_matrix())))
    header = [c.value for c in wb["Matrix"][1]]
    assert header[3:5] == ["KERUI", "KERUI documents"]
    assert wb["Matrix"]["E2"].value == "datasheet.pdf"
