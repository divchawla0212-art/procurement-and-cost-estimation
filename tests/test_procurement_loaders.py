import os
import openpyxl
import pytest
from procurement.loaders import (read_xlsx_text, read_text, read_pdf_text,
                                 read_docx_text)

FIXTURE_DOCX = os.path.join(os.path.dirname(__file__), "fixtures", "rfq_clauses.docx")


def test_read_xlsx_text(tmp_path):
    path = str(tmp_path / "q.xlsx")
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Offer"
    ws["A1"] = "Base Price"; ws["B1"] = 585000
    wb.save(path)
    text = read_xlsx_text(path)
    assert "Base Price" in text and "585000" in text and "Offer" in text


def test_read_text_dispatches_plain(tmp_path):
    p = tmp_path / "note.txt"; p.write_text("hello quote", encoding="utf-8")
    assert "hello quote" in read_text(str(p))


def test_read_pdf_text_uses_llm_fallback_when_thin(tmp_path):
    # A non-PDF byte blob: pdftotext yields nothing and pypdf raises (caught),
    # so the injected LLM fallback must be used. This proves the robustness path
    # without any network or key.
    fake = tmp_path / "scan.pdf"; fake.write_bytes(b"not a real pdf")
    called = {}

    def fb(p):
        called["path"] = p
        return "TRANSCRIBED VENDOR QUOTE TEXT " * 20

    text = read_pdf_text(str(fake), llm_fallback=fb)
    assert "TRANSCRIBED VENDOR QUOTE" in text
    assert called["path"] == str(fake)


def test_read_pdf_text_no_fallback_returns_empty_on_bad_pdf(tmp_path):
    fake = tmp_path / "scan.pdf"; fake.write_bytes(b"not a real pdf")
    assert read_pdf_text(str(fake)).strip() == ""  # never raises


def test_read_docx_text_extracts_paragraphs():
    text = read_docx_text(FIXTURE_DOCX)
    assert "MATERIAL REQUISITION MR-4471" in text
    assert "H2S tolerance shall be at least 50 ppm" in text


def test_read_docx_text_extracts_table_cells():
    """RFQ clauses live in tables as often as in prose. `Rated output` appears
    only inside the fixture's table, so a paragraph-only reader fails here."""
    text = read_docx_text(FIXTURE_DOCX)
    assert "Rated output shall be at least 1500" in text
    assert "kVA" in text
    assert "Ambient design temperature shall not exceed 50" in text
    assert "degC" in text


def test_read_docx_text_keeps_clause_and_unit_on_one_line():
    """A clause split from its unit across lines invites the model to pair the
    wrong number with the wrong unit. Cells of a row stay on one line."""
    line = [ln for ln in read_docx_text(FIXTURE_DOCX).splitlines()
            if "Rated output" in ln]
    assert len(line) == 1
    assert "5.1" in line[0] and "kVA" in line[0]


def test_read_text_dispatches_docx():
    """The bug this closes: .docx fell through to a UTF-8 raw read, and a
    deflated OOXML container decodes to mojibake that reads as real text."""
    text = read_text(FIXTURE_DOCX)
    assert "H2S tolerance shall be at least 50 ppm" in text
    assert "Rated output shall be at least 1500" in text
    # container plumbing must not reach the prompt
    assert "[Content_Types]" not in text
    assert "word/document.xml" not in text


def test_read_docx_text_raises_on_unreadable_file(tmp_path):
    """Never an empty string: "" reads downstream as "the document said
    nothing", which is exactly the silent-garbage failure being fixed."""
    bad = tmp_path / "broken.docx"
    bad.write_bytes(b"not a real docx")
    with pytest.raises(ValueError) as exc:
        read_docx_text(str(bad))
    assert "docx" in str(exc.value).lower()


def test_read_text_raises_on_unreadable_docx(tmp_path):
    bad = tmp_path / "broken.docx"
    bad.write_bytes(b"not a real docx")
    with pytest.raises(ValueError):
        read_text(str(bad))


def test_read_text_rejects_legacy_doc(tmp_path):
    """Legacy binary .doc is not an OOXML package and has no reader here. It
    must say so rather than mojibake its way into an extraction."""
    legacy = tmp_path / "old spec.doc"
    legacy.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64)
    with pytest.raises(ValueError) as exc:
        read_text(str(legacy))
    message = str(exc.value).lower()
    assert ".doc" in message
    assert "docx" in message or "pdf" in message  # says what to convert to


REAL_PDF = "data/procurement-data/ADPOWER/ADP-13158-2024-935.pdf"


@pytest.mark.skipif(not os.path.exists(REAL_PDF), reason="sample data not present")
def test_read_pdf_text_real():
    text = read_pdf_text(REAL_PDF)
    assert len(text) > 200
    assert "Baudouin" in text or "Generator" in text or "QUOTATION" in text.upper()
