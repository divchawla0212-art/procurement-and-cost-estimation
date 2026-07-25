import os
import openpyxl
import pytest
from cost_estimation.models.schema import WorkPackage
from cost_estimation.ingestion.summary_parser import (
    normalize_label, parse_summary, attach_rollups,
)


def _mini_summary(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SUMMARY"
    ws["E2"] = "Total Manhours"; ws["F2"] = "Materials"; ws["G2"] = "Consumables"
    ws["H2"] = "Installation"; ws["J2"] = "Total Value (USD)"
    ws["C3"] = "TF Main Elec. Equipment"; ws["E3"] = 85240; ws["H3"] = 2471535.06; ws["J3"] = 2471535.06
    ws["C4"] = "TF Cable and Access."; ws["J4"] = 13963935.36
    ws["J15"] = 16435470.42            # subtotal row: no label -> skipped
    ws["C16"] = "Overall SUM"; ws["J16"] = 16435470.42   # aggregate -> skipped
    wb.save(path)
    return path


def test_normalize_handles_plural_and_punctuation():
    assert normalize_label("TF Cables and Access. ") == normalize_label("TF Cable and Access.")


def test_parse_summary_extracts_packages_and_skips_aggregates(tmp_path):
    path = _mini_summary(str(tmp_path / "s.xlsx"))
    ws = openpyxl.load_workbook(path, data_only=True)["SUMMARY"]
    out = parse_summary(ws, path)
    assert set(out.keys()) == {"tf main elec equipment", "tf cable and acces"}
    assert out["tf main elec equipment"].total_value == 2471535.06
    assert out["tf main elec equipment"].installation == 2471535.06
    prov = out["tf main elec equipment"].provenance
    assert prov is not None
    assert prov.document_path == path
    assert prov.sheet == "SUMMARY"
    assert prov.cell == "C3"


def test_attach_matches_by_normalized_name(tmp_path):
    path = _mini_summary(str(tmp_path / "s.xlsx"))
    ws = openpyxl.load_workbook(path, data_only=True)["SUMMARY"]
    out = parse_summary(ws, path)
    packages = [WorkPackage(name="TF Cables and Access. ", discipline="electrical", area="TF")]
    attach_rollups(packages, out)
    assert packages[0].summary_rollup is not None
    assert packages[0].summary_rollup.total_value == 13963935.36


REAL = "data/cost-estimation-data/GDES COSTING SHEET & PROPOSAL TEMPLATE/GDX-P-26-072 REV-00(1) COSTING A-6 (Electrical BOQs for Unit areas Imp. contractor) Rev.1 RK 20260622.xlsx"


@pytest.mark.skipif(not os.path.exists(REAL), reason="sample data not present")
def test_real_summary_parses_and_reconciles_to_grand_total():
    ws = openpyxl.load_workbook(REAL, data_only=True)["SUMMARY"]
    out = parse_summary(ws, REAL)
    assert "tf main elec equipment" in out
    assert abs(out["tf main elec equipment"].total_value - 2471535.06) < 1.0
    # Sum of all parsed package rollups equals the workbook's Overall SUM row.
    grand = sum(r.total_value for r in out.values())
    assert abs(grand - 36119013.81) < 100.0
