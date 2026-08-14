"""Who Astra has approved, from a vendor-list export.

Until now the Astra overlay was **invented** — `avl_import.astra_approves`
derived it from a hash of the vendor number so a demo had something to show.
This module replaces that with a real list, read from a real file, and the
difference matters: a fabricated approval is indistinguishable on screen from a
recorded one, and this is the field the shortlist is drawn from.

The export is a subset of the client's own AVL, so it carries the same eight
columns. Only one is read here: `Vendor Number`. Everything else about a vendor
— name, product groups, manufacturers — already came from the full import, and
taking it from a second file would give two sources for one fact.

**A vendor number is digits.** The `Vendors` sheet ends with a prose footnote in
its first column, and reading that as a vendor id produced a bidder id 150
characters long. Anything non-numeric is not a vendor number and is skipped.
"""
import re

import openpyxl

# Preferred, because it is an explicit roster rather than something inferred
# from which rows happen to appear. `Sheet1` is the fallback: a plainer export
# may not carry the summary sheet at all.
_ROSTER_SHEET = "Vendors"
_VENDOR_NUMBER = "vendor number"

_IS_VENDOR_NUMBER = re.compile(r"^\d+$")


def _column(header_row, wanted: str) -> int | None:
    for index, cell in enumerate(header_row):
        if re.sub(r"\s+", " ", str(cell or "")).strip().casefold() == wanted:
            return index
    return None


def parse_vendor_numbers(path: str) -> list[str]:
    """The vendor numbers in `path`, in first-appearance order.

    Order is kept so a diff between two exports reads sensibly; duplicates are
    dropped, since a vendor listed against nine product groups is one vendor.
    """
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = (
        workbook[_ROSTER_SHEET]
        if _ROSTER_SHEET in workbook.sheetnames
        else workbook.worksheets[0]
    )
    rows = sheet.iter_rows(values_only=True)
    try:
        column = _column(next(rows), _VENDOR_NUMBER)
    except StopIteration:
        raise ValueError(f"{path} is empty — there is not even a header row.") from None
    if column is None:
        raise ValueError(
            f"This does not look like a vendor list — no column named Vendor "
            f"Number on sheet {sheet.title!r}."
        )

    seen: dict[str, None] = {}
    for row in rows:
        if column >= len(row):
            continue
        value = str(row[column] or "").strip()
        # The footnote at the bottom of the roster sheet lands here too.
        if _IS_VENDOR_NUMBER.match(value):
            seen.setdefault(value, None)
    return list(seen)
