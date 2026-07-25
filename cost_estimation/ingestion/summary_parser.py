import re
from openpyxl.utils import get_column_letter
from cost_estimation.models.schema import SummaryRollup, WorkPackage
from shared.provenance import ProvenanceRef

_VALUE_HEADERS = {
    "total_manhours": ["total manhours"],
    "materials": ["materials", "material"],
    "consumables": ["consumables"],
    "installation": ["installation"],
    "total_value": ["total value"],
}
_AGGREGATE_TOKENS = {"sum", "total", "month", "overall"}


def normalize_label(s: str) -> str:
    s = re.sub(r"[^a-z0-9 ]", " ", s.lower())
    tokens = [t[:-1] if t.endswith("s") and len(t) > 3 else t for t in s.split()]
    return " ".join(tokens)


def _num(value) -> float:
    return float(value) if isinstance(value, (int, float)) else 0.0


def _find_header_row(ws) -> int | None:
    for r in range(1, min(ws.max_row, 40) + 1):
        for c in range(1, min(ws.max_column, 40) + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and "total value" in v.lower():
                return r
    return None


def _map_value_cols(ws, header_row: int) -> dict[str, int]:
    cols: dict[str, int] = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(header_row, c).value
        if not isinstance(v, str):
            continue
        low = v.lower()
        for role, keys in _VALUE_HEADERS.items():
            if role not in cols and any(k in low for k in keys):
                cols[role] = c
    return cols


def parse_summary(worksheet, document_path: str) -> dict[str, SummaryRollup]:
    header_row = _find_header_row(worksheet)
    if header_row is None:
        return {}
    cols = _map_value_cols(worksheet, header_row)
    total_col = cols.get("total_value")
    if not total_col:
        return {}
    first_value_col = min(cols.values())
    out: dict[str, SummaryRollup] = {}
    for r in range(header_row + 1, worksheet.max_row + 1):
        total = worksheet.cell(r, total_col).value
        if not isinstance(total, (int, float)):
            continue
        label = None
        label_col = None
        for c in range(1, first_value_col):
            v = worksheet.cell(r, c).value
            if isinstance(v, str) and v.strip():
                label = v.strip()
                label_col = c
                break
        if not label:
            continue
        norm = normalize_label(label)
        if any(tok in norm.split() for tok in _AGGREGATE_TOKENS):
            continue
        out[norm] = SummaryRollup(
            total_manhours=_num(worksheet.cell(r, cols["total_manhours"]).value) if "total_manhours" in cols else 0.0,
            materials=_num(worksheet.cell(r, cols["materials"]).value) if "materials" in cols else 0.0,
            consumables=_num(worksheet.cell(r, cols["consumables"]).value) if "consumables" in cols else 0.0,
            installation=_num(worksheet.cell(r, cols["installation"]).value) if "installation" in cols else 0.0,
            total_value=float(total),
            provenance=ProvenanceRef(
                document_path=document_path,
                sheet=worksheet.title,
                cell=f"{get_column_letter(label_col)}{r}",
                extractor="xlsx",
            ),
        )
    return out


def attach_rollups(work_packages: list[WorkPackage], summary_map: dict[str, SummaryRollup]) -> None:
    for wp in work_packages:
        wp.summary_rollup = summary_map.get(normalize_label(wp.name))
