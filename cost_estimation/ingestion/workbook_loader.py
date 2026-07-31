from openpyxl.utils import get_column_letter
from pydantic import BaseModel
from shared.provenance import ProvenanceRef
from cost_estimation.models.schema import CostItem, RateBuildUp


class SheetLayout(BaseModel):
    header_row: int
    columns: dict[str, int]


def _num(value) -> float | None:
    if isinstance(value, (int, float)):
        return float(value)
    return None


def _valid_col(c) -> bool:
    return isinstance(c, int) and c >= 1


def read_items(worksheet, layout: SheetLayout, document_path: str, sheet_name: str) -> list[CostItem]:
    cols = layout.columns
    code_col = cols.get("code")
    if not _valid_col(code_col):
        code_col = None
    items: list[CostItem] = []
    for row in range(layout.header_row + 1, worksheet.max_row + 1):
        code = worksheet.cell(row, code_col).value if code_col else None
        if code is None or str(code).strip() == "":
            continue

        def cell(role):
            c = cols.get(role)
            return worksheet.cell(row, c).value if _valid_col(c) else None

        prov = ProvenanceRef(
            document_path=document_path, sheet=sheet_name,
            cell=f"{get_column_letter(code_col)}{row}", extractor="xlsx",
        )
        unit_price = _num(cell("unit_price"))
        total = _num(cell("total"))
        priced = unit_price is not None and unit_price != 0
        buildup = None
        if priced:
            buildup = RateBuildUp(
                equipment=_num(cell("equipment")) or 0.0,
                material_supply=_num(cell("material")) or 0.0,
                consumables=_num(cell("consumables")) or 0.0,
                installation=_num(cell("installation")) or 0.0,
                unit_price=unit_price,
                provenance=prov,
            )
        items.append(CostItem(
            code=str(code).strip(),
            description=str(cell("description") or "").strip(),
            uom=(str(cell("uom")).strip() if cell("uom") else None),
            quantity=_num(cell("quantity")),
            rate_buildup=buildup,
            total=total,
            priced=priced,
            provenance=prov,
        ))
    return items
