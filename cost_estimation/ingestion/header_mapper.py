from pathlib import Path
from openpyxl.utils import get_column_letter
from shared.llm.interface import LLMClient
from cost_estimation.config.loader import DisciplineConfig
from cost_estimation.ingestion.workbook_loader import SheetLayout

_PROMPT = Path(__file__).parents[2] / "shared" / "llm" / "prompts" / "header_map_v1.txt"


def sheet_preview(worksheet, rows: int = 15, cols: int = 20) -> str:
    lines = []
    for row in range(1, min(rows, worksheet.max_row) + 1):
        cells = []
        for col in range(1, min(cols, worksheet.max_column) + 1):
            value = worksheet.cell(row, col).value
            if value is not None:
                cells.append(f"{get_column_letter(col)}={value}")
        if cells:
            lines.append(f"R{row} | " + " | ".join(cells))
    return "\n".join(lines)


def map_sheet(worksheet, client: LLMClient, config: DisciplineConfig, sheet_name: str) -> SheetLayout:
    prompt = _PROMPT.read_text(encoding="utf-8")
    preview = sheet_preview(worksheet)
    result = client.classify_structure(prompt, SheetLayout, preview)
    return SheetLayout.model_validate(result)
