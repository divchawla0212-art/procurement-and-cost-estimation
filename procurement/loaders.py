import os
import shutil
import subprocess
import docx
import openpyxl
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl.utils import get_column_letter


def read_xlsx_text(path: str) -> str:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    lines: list[str] = []
    for ws in wb.worksheets:
        for r, row in enumerate(ws.iter_rows(values_only=True), 1):
            cells = [
                f"{get_column_letter(c)}{r}={v}"
                for c, v in enumerate(row, 1)
                if v is not None
            ]
            if cells:
                lines.append(f"{ws.title} | " + " | ".join(cells))
    return "\n".join(lines)


def _docx_lines(parent, element) -> list[str]:
    """Body text of one container, in document order.

    python-docx exposes `.paragraphs` and `.tables` as separate flat lists, so
    reading them in turn would tear a numbered clause away from the table that
    qualifies it. Walking the XML children keeps the original order.
    """
    lines: list[str] = []
    for child in element.iterchildren():
        if child.tag == qn("w:p"):
            text = Paragraph(child, parent).text.strip()
            if text:
                lines.append(text)
        elif child.tag == qn("w:tbl"):
            for row in Table(child, parent).rows:
                cells: list[str] = []
                seen: set[int] = set()
                for cell in row.cells:
                    # A horizontally merged cell is yielded once per column it
                    # spans; emitting it each time would triple a clause.
                    if id(cell._tc) in seen:
                        continue
                    seen.add(id(cell._tc))
                    # Nested tables collapse into their host cell rather than
                    # becoming rows of their own.
                    cells.append(" ".join(_docx_lines(cell, cell._tc)).strip())
                if any(cells):
                    # One row per line, cells joined as read_xlsx_text joins
                    # them: a clause split from its unit across lines invites
                    # the model to pair the wrong number with the wrong unit.
                    lines.append(" | ".join(cells))
    return lines


def read_docx_text(path: str) -> str:
    """Paragraph and table text of a Word (.docx) document.

    Raises ValueError if the file is not a readable OOXML package. It must not
    degrade to "": an empty string reads downstream as "the document said
    nothing", and a requirements extraction would be recorded `ok` with no
    clauses rather than `failed` with a reason.
    """
    try:
        document = docx.Document(path)
        return "\n".join(_docx_lines(document, document.element.body))
    except Exception as exc:
        raise ValueError(f"unreadable .docx document {path!r}: {exc}") from exc


def _pdftotext(path: str) -> str:
    exe = shutil.which("pdftotext")
    if not exe:
        return ""
    try:
        out = subprocess.run([exe, "-layout", path, "-"], capture_output=True, timeout=120)
        return out.stdout.decode("utf-8", errors="ignore")
    except Exception:
        return ""


def _pypdf(path: str) -> str:
    try:
        from pypdf import PdfReader
        reader = PdfReader(path)
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    except Exception:
        return ""


# Below this many characters of readable text, a document is not extracted.
# Chosen against the live corpus: the five KERUI drawings yield 1-42 chars,
# and "11 Attachment-4 Applicable Codes and Standards Pending.pdf" yields 168
# and is a real, if short, document. Illustrative, not load-bearing — but the
# 1-42 range must fail the guard and 168 must not, so test fixtures standing
# in for extractable documents carry realistic (>=100 char) bodies rather
# than pulling this threshold down to fit short synthetic text.
MIN_EXTRACTABLE_CHARS = 100


def _llamaparse(path: str) -> str:
    # Imported here, not at module scope: llama-parse is an optional extra, and
    # `loaders` is on the import path of the whole pipeline. A missing optional
    # dependency must not stop anyone reading a .docx.
    from procurement.llamaparse_reader import parse_pdf
    return parse_pdf(path)


# One reader, no chain. Named arms for the parser A/B; production passes None
# and keeps the fallback chain below. The lambdas resolve the module globals at
# call time so a test can stub `_pdftotext` and have the arm see the stub.
_PURE_PDF_READERS = {
    "pdftotext": lambda path: _pdftotext(path),
    "pypdf": lambda path: _pypdf(path),
    "llamaparse": lambda path: _llamaparse(path),
}


def read_pdf_text_with_source(path: str, llm_fallback=None,
                              min_chars: int = 200,
                              pdf_reader: str | None = None) -> tuple[str, str]:
    """Robust PDF text plus the reader that produced it: pdftotext CLI ->
    pypdf -> LLM transcription fallback.

    `pdf_reader` names a single reader and disables the chain entirely,
    including the LLM fallback. It exists for the parser comparison, where a
    fallback would be actively misleading: an arm that quietly falls through
    reports another reader's numbers under its own name, and because a rescue
    fires *least* for the best parser, every arm would converge and the
    experiment would read "no difference" whatever the truth. Production passes
    None and is unaffected.

    A pure arm still returns "" for a document it cannot read. That is a result
    -- the caller's MIN_EXTRACTABLE_CHARS guard records it as `failed` with the
    reader's name and the character count -- so it must not be rescued here.
    """
    if pdf_reader is not None:
        reader = _PURE_PDF_READERS.get(pdf_reader)
        if reader is None:
            raise ValueError(
                f"unknown pdf_reader {pdf_reader!r}; "
                f"expected one of {sorted(_PURE_PDF_READERS)} or None")
        return reader(path), pdf_reader

    text, source = _pdftotext(path), "pdftotext"
    if len(text.strip()) < min_chars:
        alt = _pypdf(path)
        if len(alt.strip()) > len(text.strip()):
            text, source = alt, "pypdf"
    if len(text.strip()) < min_chars and llm_fallback is not None:
        try:
            fb = llm_fallback(path)
            if fb and fb.strip():
                return fb, "llm"
        except Exception:
            pass
    return text, source


def read_pdf_text(path: str, llm_fallback=None, min_chars: int = 200,
                  pdf_reader: str | None = None) -> str:
    return read_pdf_text_with_source(path, llm_fallback, min_chars,
                                     pdf_reader=pdf_reader)[0]


def read_text_with_source(path: str, llm_fallback=None,
                          pdf_reader: str | None = None) -> tuple[str, str]:
    """Text plus the name of the reader that produced it.

    `text_source` has been declared on DocumentRecord since phase 2 and never
    populated. It is the only way to tell "this datasheet genuinely states
    nothing" from "we read this scan with the wrong reader".
    """
    ext = os.path.splitext(path)[1].lower()
    if ext == ".xlsx":
        return read_xlsx_text(path), "xlsx"
    if ext == ".pdf":
        # `pdf_reader` is deliberately consulted only here. The A/B varies the
        # PDF reader; every other format keeps its single reader in every arm,
        # which is what makes the .xlsx control a control.
        return read_pdf_text_with_source(path, llm_fallback=llm_fallback,
                                         pdf_reader=pdf_reader)
    if ext == ".docx":
        return read_docx_text(path), "docx"
    if ext == ".doc":
        # Legacy binary Word, not an OOXML package — nothing here reads it. The
        # raw fallback below would decode it to mojibake and hand that to an
        # extractor as though it were clause text, so say so instead.
        raise ValueError(
            f"legacy binary .doc is not supported: {path!r}. "
            "Convert it to .docx or .pdf and re-upload.")
    with open(path, "r", encoding="utf-8", errors="ignore") as fh:
        return fh.read(), "text"


def read_text(path: str, llm_fallback=None,
              pdf_reader: str | None = None) -> str:
    return read_text_with_source(path, llm_fallback=llm_fallback,
                                 pdf_reader=pdf_reader)[0]
