import io
import csv
import openpyxl
from procurement.models import ComparisonTable
from procurement.matrix import ComplianceMatrix, bound

_HEADERS = ["vendor", "currency", "raw_base_price", "normalized_total",
            "delivery_terms", "delivery_time", "payment_terms", "engine_make",
            "extraction_status", "normalization_status"]


def comparison_to_rows(table: ComparisonTable) -> list[dict]:
    return [{h: getattr(r, h) for h in _HEADERS} for r in table.rows]


def comparison_to_csv_str(table: ComparisonTable) -> str:
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=_HEADERS)
    writer.writeheader()
    for row in comparison_to_rows(table):
        writer.writerow(row)
    return buf.getvalue()


def comparison_to_xlsx_bytes(table: ComparisonTable) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Comparison"
    ws.append(_HEADERS)
    for row in comparison_to_rows(table):
        ws.append([row[h] for h in _HEADERS])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _statement_value(cell) -> str:
    """One rule, one place. The note is composed with the value, never chosen
    against it: `cell.text or cell.note` would drop the compliance tally for
    every vendor that also has a reviewer's note — the exact pairing the
    statement exists to show — and would strip FINAL VALUE's "base price not
    stated" off a blank cell, leaving an empty total with no reason."""
    if cell is None:
        return ""
    parts = []
    if cell.total is not None:
        parts.append(f"{cell.total:,.2f}")
    parts.extend(p for p in (cell.text, cell.note) if p)
    return " — ".join(parts)


def statement_to_rows(statement) -> list[dict]:
    """Flat rows for CSV/XLSX: Description plus one column per vendor.

    The three Qty/Unit/Total sub-columns collapse to the total here — the
    export is a record of the comparison, and a spreadsheet reader can open
    the source quotation for the breakdown.
    """
    rows = []
    for row in statement.rows:
        if not row.cells:
            continue
        out = {"Description": row.label}
        for vendor in statement.vendors:
            out[vendor] = _statement_value(row.cells.get(vendor))
        rows.append(out)
    return rows


def statement_to_csv_str(statement) -> str:
    rows = statement_to_rows(statement)
    if not rows:
        return ""
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=list(rows[0].keys()))
    writer.writeheader()
    writer.writerows(rows)
    return buf.getvalue()


def statement_to_xlsx_bytes(statement) -> bytes:
    rows = statement_to_rows(statement)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Comparative Statement"
    if rows:
        headers = list(rows[0].keys())
        ws.append(headers)
        for row in rows:
            ws.append([row[h] for h in headers])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------ compliance matrix
#
# Two shapes, because one cannot serve both readers. The grid is what a
# procurement reviewer opens — clause down the side, vendors across, one word
# per cell — and inlining the rationale there would make every row as tall as
# its longest sentence. The detail sheet is the audit trail: the verdict beside
# the reason it was reached and the evidence it was reached from. A verdict
# without its rationale is not actionable away from the screen, so the export
# carries both rather than choosing.

_GRID_PREFIX = ["Clause", "Requirement", "Checkability"]
_DETAIL_PREFIX = ["Clause", "Requirement", "Clause text", "Checkability"]
_DETAIL_CELL = ["Vendor", "Verdict", "Group", "Rationale", "Fact id", "Doc id",
                "Doc name", "Candidate fact ids", "Candidate documents"]
_DETAIL_HEADERS = _DETAIL_PREFIX + _DETAIL_CELL


def _documents_column(vendor: str) -> str:
    return f"{vendor} documents"


def grid_headers(matrix: ComplianceMatrix) -> list[str]:
    """The grid's columns, in order. One source for the row dicts, the CSV
    field list and the worksheet header — three copies of this interleaving is
    how a column silently shifts one place against its heading."""
    columns = list(_GRID_PREFIX)
    for vendor in matrix.vendors:
        columns += [vendor, _documents_column(vendor)]
    return columns


def _cell_documents(cell) -> str:
    """The file(s) the verdict was read from, as a reviewer would name them.

    Candidates first: when a cell was decided against more than one reading,
    every document behind those readings is material — that multiplicity is
    usually *why* the cell is a `review`. `doc_name` is the fallback for the
    ordinary single-reading cell, which carries no candidate list at all.
    """
    if cell is None:
        return ""
    if cell.candidate_doc_names:
        return "; ".join(cell.candidate_doc_names)
    return cell.doc_name or ""


def matrix_to_grid_rows(matrix: ComplianceMatrix) -> list[dict]:
    """One row per requirement; per vendor, the verdict and the document(s) it
    was read from.

    A vendor with no computed cell exports blank. Never `pass`, never `0`, never
    a dash that a spreadsheet reader could mistake for a value: the pipeline not
    having produced a cell is not the vendor having answered. A blank documents
    cell is the same refusal — a judgement row and a silent `unanswered` cite no
    document, and naming one would credit a file that produced no verdict.
    """
    rows = []
    for row in matrix.rows:
        out = {"Clause": row.clause_ref, "Requirement": bound(row),
               "Checkability": row.checkability}
        for vendor in matrix.vendors:
            cell = row.cells.get(vendor)
            out[vendor] = cell.verdict if cell else ""
            out[_documents_column(vendor)] = _cell_documents(cell)
        rows.append(out)
    return rows


def matrix_to_detail_rows(matrix: ComplianceMatrix) -> list[dict]:
    """One row per (requirement, vendor), carrying the rationale and evidence.

    `Clause text` is kept alongside `Requirement` because `bound` replaces the
    prose with the checkable bound on every `auto` and `stated` row — without
    this column the reviewer's own wording never reaches the export.

    The rationale is copied verbatim and never parsed (compliance spec §5).
    """
    rows = []
    for row in matrix.rows:
        for vendor in matrix.vendors:
            cell = row.cells.get(vendor)
            rows.append({
                "Clause": row.clause_ref,
                "Requirement": bound(row),
                "Clause text": row.text,
                "Checkability": row.checkability,
                "Vendor": vendor,
                "Verdict": cell.verdict if cell else "",
                "Group": cell.group if cell else "",
                "Rationale": cell.rationale if cell else "",
                "Fact id": (cell.fact_id or "") if cell else "",
                "Doc id": (cell.doc_id or "") if cell else "",
                # Kept beside the id, not instead of it: two files sharing a
                # basename inside one vendor folder read alike by name, and the
                # id is what disambiguates them.
                "Doc name": (cell.doc_name or "") if cell else "",
                # Joined rather than one column per candidate: the count varies
                # per cell, and a ragged header is not a spreadsheet.
                "Candidate fact ids": ", ".join(cell.candidate_fact_ids) if cell else "",
                "Candidate documents": ", ".join(cell.candidate_doc_names) if cell else "",
            })
    return rows


def matrix_to_csv_str(matrix: ComplianceMatrix) -> str:
    """The grid alone. CSV is one table, and of the two the grid is the one a
    reviewer scans; the detail sheet is why the Excel export exists."""
    rows = matrix_to_grid_rows(matrix)
    if not rows:
        return ""
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=grid_headers(matrix))
    writer.writeheader()
    writer.writerows(rows)
    return buf.getvalue()


def matrix_to_xlsx_bytes(matrix: ComplianceMatrix) -> bytes:
    wb = openpyxl.Workbook()
    grid_ws = wb.active
    grid_ws.title = "Matrix"
    # Headers are written even for an empty matrix: a workbook whose sheets are
    # blank reads as a failed export rather than as an empty project.
    headers = grid_headers(matrix)
    grid_ws.append(headers)
    for row in matrix_to_grid_rows(matrix):
        grid_ws.append([row[h] for h in headers])

    detail_ws = wb.create_sheet("Detail")
    detail_ws.append(_DETAIL_HEADERS)
    for row in matrix_to_detail_rows(matrix):
        detail_ws.append([row[h] for h in _DETAIL_HEADERS])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
